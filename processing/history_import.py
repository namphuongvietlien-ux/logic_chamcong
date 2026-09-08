"""Parse finalized Chi_Tiet_Cham_Cong workbooks (sheet Tổng hợp) for historical import."""

from __future__ import annotations

import calendar
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from processing.database import (
    find_employee_by_code,
    find_employee_by_name_key,
    list_employees,
    lock_month,
    month_key,
    parse_month_key,
    replace_attendance_history,
    replace_attendance_month,
)
from processing.excel_locale import excel_formula
from processing.leave import LEAVE_TYPE_PAID, create_leave_request, invalidate_leave_cache
from processing.period import period_from_text
from processing.utils import name_match_key

SHEET_NAME = "Tổng hợp"
HISTORY_NOTE_PREFIX = "Nạp lịch sử "

HEADER_FIELDS = {
    "employee_id": (
        "mã nv",
        "ma nv",
        "mã nhân viên",
        "ma nhan vien",
        "employee id",
        "employee_id",
        "mã",
    ),
    "employee_name": (
        "tên nhân viên",
        "ten nhan vien",
        "họ và tên",
        "ho va ten",
        "họ tên",
        "ho ten",
        "employee name",
        "tên nv",
        "ten nv",
    ),
    "workdays": (
        "tổng công",
        "tong cong",
        "công",
        "cong",
        "standardized workday",
        "total workdays",
    ),
    "used_leave": (
        "phép năm",
        "phep nam",
        "phép đã nghỉ",
        "phep da nghi",
        "used leave",
        "phép",
    ),
    "overtime": (
        "tăng ca",
        "tang ca",
        "overtime",
        "ot",
        "giờ tăng ca",
    ),
}


def _fold(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return " ".join(text.replace("_", " ").split())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _num(value: Any) -> float:
    if value is None:
        return 0.0
    try:
        if pd.isna(value):
            return 0.0
    except (TypeError, ValueError):
        pass
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def guess_month_year(path: str | Path, sheet_hint: str = "") -> str:
    parsed = period_from_text(f"{Path(path).name} {sheet_hint}")
    if parsed:
        year, month = parsed
        return month_key(year, month)
    today = date.today()
    return month_key(today.year, today.month)


def resolve_tong_hop_sheet(path: str | Path) -> str:
    wanted = _fold(SHEET_NAME)
    with pd.ExcelFile(path) as xl:
        for name in xl.sheet_names:
            if _fold(name) == wanted:
                return name
    raise ValueError(f'Không tìm thấy sheet "{SHEET_NAME}" trong file {Path(path).name}.')


def _match_field(label: str) -> str | None:
    folded = _fold(label)
    if not folded:
        return None
    for field, aliases in HEADER_FIELDS.items():
        for alias in aliases:
            if folded == alias or folded.startswith(alias + " ") or alias in folded:
                return field
    return None


def _score_header(cells: list[str]) -> int:
    found: set[str] = set()
    for cell in cells:
        field = _match_field(cell)
        if field:
            found.add(field)
    if "workdays" not in found:
        return 0
    if "employee_id" not in found and "employee_name" not in found:
        return 0
    return len(found)


def _find_header_row(raw: pd.DataFrame) -> tuple[int, dict[str, int]] | None:
    limit = min(len(raw), 60)
    best: tuple[int, dict[str, int], int] | None = None
    for index in range(limit):
        cells = [_cell_text(v) for v in raw.iloc[index].tolist()]
        score = _score_header(cells)
        if score < 2:
            continue
        mapping: dict[str, int] = {}
        for col, cell in enumerate(cells):
            field = _match_field(cell)
            if field and field not in mapping:
                mapping[field] = col
        if score > (best[2] if best else 0):
            best = (index, mapping, score)
    if best is None:
        return None
    return best[0], best[1]


def _parse_tabular(raw: pd.DataFrame) -> list[dict]:
    found = _find_header_row(raw)
    if found is None:
        return []
    header_idx, mapping = found
    records: list[dict] = []
    for index in range(header_idx + 1, len(raw)):
        row = raw.iloc[index]
        emp_id = _cell_text(row.iloc[mapping["employee_id"]]) if "employee_id" in mapping else ""
        name = _cell_text(row.iloc[mapping["employee_name"]]) if "employee_name" in mapping else ""
        if not emp_id and not name:
            continue
        blob = " ".join(_cell_text(v) for v in row.tolist())
        if "bảng chi tiết" in _fold(blob) or blob.upper().startswith("TỔNG CỘNG"):
            continue
        if _score_header([_cell_text(v) for v in row.tolist()]) >= 2:
            continue
        records.append(
            {
                "employee_id": emp_id.replace("-----", "").strip(),
                "employee_name": name,
                "workdays": _num(row.iloc[mapping["workdays"]]) if "workdays" in mapping else 0.0,
                "used_leave": _num(row.iloc[mapping["used_leave"]]) if "used_leave" in mapping else 0.0,
                "overtime_hours": _num(row.iloc[mapping["overtime"]]) if "overtime" in mapping else 0.0,
            }
        )
    return records


def _split_labeled(text: str) -> tuple[str, str]:
    if ":" not in text:
        return _fold(text), ""
    label, value = text.split(":", 1)
    return _fold(label), value.strip()


def _parse_blocks(raw: pd.DataFrame) -> list[dict]:
    """Fallback for this app's stacked Chi tiết / Tổng hợp layout."""
    records: list[dict] = []
    nrows, ncols = raw.shape
    i = 0
    while i < nrows:
        emp_id = ""
        name = ""
        for col in range(ncols):
            text = _cell_text(raw.iat[i, col])
            label, value = _split_labeled(text)
            if label.startswith("ma nhan vien") or label == "ma nv":
                emp_id = value or emp_id
            elif label.startswith("ten nhan vien") or label.startswith("ho va ten"):
                name = value or name
        if emp_id or name:
            workdays = used_leave = overtime = 0.0
            for look in range(i, min(i + 10, nrows)):
                for col in range(ncols):
                    label = _fold(_cell_text(raw.iat[look, col]))
                    nxt = raw.iat[look, col + 1] if col + 1 < ncols else None
                    if label == "tong cong":
                        workdays = _num(nxt)
                    elif label in {"tang ca", "gio tang ca"}:
                        overtime = _num(nxt)
                    elif label in {"phep nam", "phep da nghi"}:
                        used_leave = _num(nxt)
            records.append(
                {
                    "employee_id": emp_id.replace("-----", "").strip(),
                    "employee_name": name,
                    "workdays": workdays,
                    "used_leave": used_leave,
                    "overtime_hours": overtime,
                }
            )
        i += 1
    return records


def extract_tong_hop(path: str | Path) -> dict[str, Any]:
    """Read sheet Tổng hợp and return preview rows (not yet saved)."""
    file = Path(path)
    sheet = resolve_tong_hop_sheet(file)
    raw = pd.read_excel(file, sheet_name=sheet, header=None, dtype=object)
    if raw.empty:
        raise ValueError(f'Sheet "{sheet}" trống.')
    records = _parse_tabular(raw)
    source = "header"
    if len(records) < 1:
        records = _parse_blocks(raw)
        source = "block"
    if not records:
        raise ValueError(
            'Không tìm thấy dòng tiêu đề chứa "Mã NV"/"Tên nhân viên", "Tổng công", '
            '"Phép năm" hoặc "Tăng ca" trên sheet Tổng hợp.'
        )
    guessed = guess_month_year(file, sheet)
    return {
        "path": str(file),
        "sheet": sheet,
        "month_year": guessed,
        "parse_mode": source,
        "rows": records,
    }


def match_history_rows(rows: list[dict]) -> list[dict]:
    out: list[dict] = []
    for item in rows:
        code = str(item.get("employee_id") or "").strip()
        name = str(item.get("employee_name") or "").strip()
        emp = find_employee_by_code(code) if code else None
        if emp is None and name:
            emp = find_employee_by_name_key(name_match_key(name))
        matched = dict(item)
        if emp:
            matched["matched"] = True
            matched["employee_pk"] = int(emp["id"])
            matched["name_key"] = emp.get("name_key") or name_match_key(emp.get("employee_name") or name)
            matched["employee_id"] = emp.get("employee_id") or code
            matched["employee_name"] = emp.get("employee_name") or name
            matched["match_note"] = "Khớp"
        else:
            matched["matched"] = False
            matched["employee_pk"] = None
            matched["name_key"] = name_match_key(name)
            matched["match_note"] = "Không có trong CSDL"
        out.append(matched)
    return out


def history_leave_note(month_year: str) -> str:
    return f"{HISTORY_NOTE_PREFIX}{month_year}"


def _delete_history_leave(month_year: str) -> None:
    from processing.database import connect, init_db

    init_db()
    note = history_leave_note(month_year)
    with connect() as conn:
        conn.execute("DELETE FROM leave_requests WHERE note = ?", (note,))
    invalidate_leave_cache()


def write_missing_employee_log(unmatched: list[dict], dest: str | Path) -> Path:
    """Error workbook for HR; VLOOKUP uses ';' for VN/EU Excel."""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    missing = wb.active
    missing.title = "Không khớp"
    headers = ("Mã NV", "Tên nhân viên", "Tổng công", "Phép năm", "Tăng ca", "Ghi chú", "Tên từ CSDL (VLOOKUP)")
    missing.append(headers)
    for cell in missing[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C53030")
    for index, item in enumerate(unmatched, start=2):
        missing.cell(index, 1, item.get("employee_id") or "")
        missing.cell(index, 2, item.get("employee_name") or "")
        missing.cell(index, 3, float(item.get("workdays") or 0))
        missing.cell(index, 4, float(item.get("used_leave") or 0))
        missing.cell(index, 5, float(item.get("overtime_hours") or 0))
        missing.cell(index, 6, item.get("match_note") or "Không có trong CSDL")
        missing.cell(index, 7, excel_formula("VLOOKUP", f"A{index}", "DanhSachNV!A:B", "2", "0"))

    roster = wb.create_sheet("DanhSachNV")
    roster.append(("Mã NV", "Tên nhân viên"))
    for cell in roster[1]:
        cell.font = Font(bold=True)
    for emp in list_employees(active_only=False):
        roster.append((emp.get("employee_id") or "", emp.get("employee_name") or ""))

    wb.save(dest)
    wb.close()
    return dest


def commit_history_import(
    rows: list[dict],
    month_year: str,
    source_file: str = "",
    error_dir: str | Path | None = None,
) -> dict[str, Any]:
    """Save matched aggregates, lock the month, deduct paid leave. Write a miss log if needed."""
    matched = [row for row in rows if row.get("matched") and row.get("employee_pk")]
    unmatched = [row for row in rows if not row.get("matched")]
    if not matched:
        raise ValueError("Không có nhân viên nào khớp CSDL để lưu lịch sử.")
    year, month = parse_month_key(month_year)
    first_day = date(year, month, 1)
    last_day = date(year, month, calendar.monthrange(year, month)[1])

    history_rows = []
    day_rows = []
    for row in matched:
        history_rows.append(
            {
                "name_key": row.get("name_key"),
                "employee_id": row.get("employee_id") or "",
                "employee_name": row.get("employee_name") or "",
                "workdays": float(row.get("workdays") or 0),
                "used_leave": float(row.get("used_leave") or 0),
                "overtime_hours": float(row.get("overtime_hours") or 0),
            }
        )
        day_rows.append(
            {
                "name_key": row.get("name_key"),
                "employee_name": row.get("employee_name") or "",
                "work_date": first_day,
                "standardized_workday": float(row.get("workdays") or 0),
                "overtime_hours": float(row.get("overtime_hours") or 0),
                "notes": f"{history_leave_note(month_year)} · sheet Tổng hợp",
            }
        )

    replace_attendance_history(month_year, history_rows, source_file=Path(source_file).name if source_file else "")
    replace_attendance_month(month_year, day_rows, ignore_lock=True)

    _delete_history_leave(month_year)
    leave_count = 0
    for row in matched:
        used = float(row.get("used_leave") or 0)
        if used <= 0:
            continue
        create_leave_request(
            {
                "employee_id": int(row["employee_pk"]),
                "leave_type": LEAVE_TYPE_PAID,
                "start_date": first_day.isoformat(),
                "end_date": last_day.isoformat(),
                "days": used,
                "session": "Cả ngày",
                "status": "approved",
                "note": history_leave_note(month_year),
            }
        )
        leave_count += 1
    invalidate_leave_cache()
    lock_month(month_year, locked_by="Nạp lịch sử")

    miss_path: Optional[Path] = None
    if unmatched:
        folder = Path(error_dir) if error_dir else Path(source_file).parent if source_file else Path.cwd()
        stem = Path(source_file).stem if source_file else f"Chi_Tiet_{month_year.replace('/', '.')}"
        miss_path = write_missing_employee_log(unmatched, folder / f"{stem}_khong_khop.xlsx")

    return {
        "saved": len(matched),
        "unmatched": len(unmatched),
        "leave_requests": leave_count,
        "month_year": month_year,
        "miss_log": miss_path,
    }

"""Load K9 / working table, flag missing punches, apply manual clock edits, re-export reports."""

from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
from openpyxl import load_workbook

from processing.cong_rules import (
    _as_time,
    filled_clocks,
    hours_from_clocks,
    workday_count,
)
from processing.master_data import lookup_config
from processing.leave import REMARK_PAID, REMARK_PAID_HALF, REMARK_UNPAID, REMARK_UNPAID_HALF
from processing.merger import recompute_record
from processing.sessions import (
    STATUS_MISSING_IN,
    STATUS_MISSING_OUT,
    STATUS_ODD_PUNCH,
    clocks_to_datetimes,
    four_clocks_dt,
    merge_day_punches,
    missing_punch_status,
)
from processing.period import period_from_text
from processing.utils import (
    clock_plus_days,
    employee_id_sort_key,
    format_time,
    name_match_key,
    parse_date_value,
    parse_time_value,
    sort_frame_by_employee_id,
)

LogFn = Optional[Callable[[str], None]]

SHEET_DAY_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})$")
WORKING_COLUMNS = [
    "employee_id",
    "employee_name",
    "department",
    "date",
    "in1",
    "out1",
    "in2",
    "out2",
    "in1_plus",
    "out1_plus",
    "in2_plus",
    "out2_plus",
    "overnight",
    "standard_shift_hours",
    "shift_start",
    "shift_end",
    "lunch_duration_hours",
    "actual_work_hours",
    "deducted_lunch_hours",
    "standardized_workday",
    "overtime_hours",
    "late_minutes",
    "early_minutes",
    "notes",
    "manual_edit",
    "is_leave",
    "leave_type",
    "leave_duration",
    "leave_session",
]


def shift_target_hours(standard) -> float:
    """Employee shift length used as the 8h / 12h milestone (from master)."""
    try:
        hours = float(standard if standard is not None and str(standard).strip() != "" else 8)
    except (TypeError, ValueError):
        return 8.0
    if hours <= 0:
        return 8.0
    return hours


def _hour_label(value: float) -> str:
    if abs(value - round(value)) < 0.05:
        return str(int(round(value)))
    return f"{value:.1f}"


def milestone_from_clocks(
    in1,
    out1,
    in2,
    out2,
    standard_shift_hours=None,
    lunch_duration_hours=None,
    overnight=False,
    work_date=None,
) -> dict:
    """Hours vs 8h or 12h target so clerks see 8/8 ĐẠT vs 8/12 CHƯA ĐẠT.

    Example: ca 8h vào 12:00 ra 20:00 → 8/8 ĐẠT (không cần sửa).
    Ca 12h cùng giờ → 8/12 CHƯA ĐẠT (cần chỉnh).
    """
    in1, out1, in2, out2 = _as_time(in1), _as_time(out1), _as_time(in2), _as_time(out2)
    target = shift_target_hours(standard_shift_hours)
    hours, _deducted = hours_from_clocks(
        in1, out1, in2, out2, lunch_duration_hours, logical_date=work_date, overnight=bool(overnight)
    )
    cong = workday_count(hours, target)
    actual = float(hours or 0)
    full_mark = target * 0.85
    reached = actual + 0.1 >= full_mark
    fraction = f"{_hour_label(actual)}/{_hour_label(target)}"
    hours_text = f"{_hour_label(actual)}/{_hour_label(target)}h"
    badge = "ĐẠT" if reached else "CHƯA ĐẠT"
    if cong >= 1:
        cong_text = "1 công"
    elif cong >= 0.5:
        cong_text = "0.5 công"
    else:
        cong_text = "0 công"
    return {
        "hours": hours,
        "actual": actual,
        "target": target,
        "reached": reached,
        "fraction": fraction,
        "hours_text": hours_text,
        "badge": badge,
        "cong": cong,
        "cong_text": cong_text,
        "shift_text": f"Ca {_hour_label(target)}h",
        "worked_text": f"{_hour_label(actual)}h làm",
    }


def milestone_from_row(row) -> dict:
    return milestone_from_clocks(
        row.get("in1"),
        row.get("out1"),
        row.get("in2"),
        row.get("out2"),
        row.get("standard_shift_hours"),
        row.get("lunch_duration_hours"),
        overnight=bool(row.get("overnight")),
        work_date=row.get("date"),
    )


def _leave_flags(row) -> tuple[bool, bool]:
    """(authorized_full_day, half_day) from computed fields or Ghi chú text."""
    notes = str(row.get("notes") or "")
    duration = 0.0
    try:
        duration = float(row.get("leave_duration") or 0)
    except (TypeError, ValueError):
        duration = 0.0
    half = bool(row.get("is_leave")) and 0 < duration < 1
    half = half or REMARK_PAID_HALF in notes or REMARK_UNPAID_HALF in notes or "nửa ngày" in notes
    full = bool(row.get("is_leave")) and duration >= 1
    full = full or (REMARK_PAID in notes and REMARK_PAID_HALF not in notes)
    full = full or (REMARK_UNPAID in notes and REMARK_UNPAID_HALF not in notes)
    return full, half


def needs_review(row) -> bool:
    """Missing in/out/lunch, or hours short of the 8h / 12h shift."""
    full, half = _leave_flags(row)
    issue = str(row.get("missing_punch") or "").strip() or issue_label(
        row.get("in1"),
        row.get("out1"),
        row.get("in2"),
        row.get("out2"),
        overnight=bool(row.get("overnight")),
        authorized_leave=full,
        half_day_leave=half,
    )
    if issue:
        return True
    if full:
        return False
    if half:
        return False
    return not bool(milestone_from_row(row)["reached"])


def issue_label(in1, out1, in2, out2, overnight=False, authorized_leave=False, half_day_leave=False) -> str:
    """Human-readable missing-punch reason. Empty if the day looks complete."""
    if authorized_leave:
        return ""
    slots = (_as_time(in1), _as_time(out1), _as_time(in2), _as_time(out2))
    n = filled_clocks(*slots)
    in1, out1, in2, out2 = slots
    last = out2 or in2 or out1
    if half_day_leave:
        if n == 0:
            return ""
        if (in1 and (out1 or out2)) or (in2 and out2):
            return ""
    if overnight and (in1 or in2) and (out1 or out2):
        if n == 3 and in2 is None:
            return STATUS_ODD_PUNCH
        return ""
    if in2 is not None and out2 is None:
        return "Thiếu giờ ra"
    if in1 is not None and out1 is None and out2 is None:
        return "Thiếu giờ ra"
    if n == 0:
        return ""
    if n == 1:
        punch = in1 or out1 or in2 or out2
        if punch is not None and punch.hour >= 12:
            return STATUS_MISSING_IN
        return STATUS_MISSING_OUT
    if n == 2:
        if in1 and out2 and not out1 and not in2:
            return ""
        if in1 and out1 and not in2 and not out2:
            return ""
        if in1 and not last:
            return "Thiếu giờ ra"
        if last and not in1:
            return "Thiếu giờ vào"
        return "Thiếu mốc chấm"
    if n == 3:
        return STATUS_ODD_PUNCH
    return ""


def _cell_time(value) -> Optional[time]:
    return parse_time_value(value)


def _as_work_date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date_value(value)


def is_k9_workbook(path: str | Path) -> bool:
    wb = load_workbook(path, read_only=True)
    try:
        return any(SHEET_DAY_RE.match(str(name).strip()) for name in wb.sheetnames)
    finally:
        wb.close()


def load_k9_file(path: str | Path, master: pd.DataFrame | None = None, log: LogFn = None) -> tuple[pd.DataFrame, int, int]:
    """Build a working table from K9_MM.YYYY.xlsx (one sheet per day)."""
    file = Path(path)
    pair = period_from_text(file.stem) or period_from_text(file.name)
    wb = load_workbook(file, data_only=True)
    try:
        rows: list[dict] = []
        roster: list[str] = []
        year = pair[0] if pair else None
        month = pair[1] if pair else None
        for sheet_name in wb.sheetnames:
            match = SHEET_DAY_RE.match(str(sheet_name).strip())
            if not match:
                continue
            day_n, month_n = int(match.group(1)), int(match.group(2))
            if month is None:
                month = month_n
            if year is None:
                year = date.today().year
            work_date = date(int(year), month_n, day_n)
            ws = wb[sheet_name]
            for r in range(5, (ws.max_row or 5) + 1):
                name = ws.cell(r, 3).value
                if not name or not str(name).strip():
                    continue
                name = str(name).strip()
                if name not in roster:
                    roster.append(name)
                in1 = _cell_time(ws.cell(r, 4).value)
                out1 = _cell_time(ws.cell(r, 5).value)
                in2 = _cell_time(ws.cell(r, 6).value)
                out2 = _cell_time(ws.cell(r, 7).value)
                plus = max(
                    clock_plus_days(ws.cell(r, 4).value),
                    clock_plus_days(ws.cell(r, 5).value),
                    clock_plus_days(ws.cell(r, 6).value),
                    clock_plus_days(ws.cell(r, 7).value),
                )
                if not any((in1, out1, in2, out2)):
                    continue
                cfg = lookup_config(master if master is not None else pd.DataFrame(), name)
                data = {
                    "employee_name": name,
                    "date": work_date,
                    "in1": in1,
                    "out1": out1,
                    "in2": in2,
                    "out2": out2,
                    "overnight": plus > 0,
                    "manual_edit": plus > 0,
                    **cfg,
                }
                rows.append(recompute_record(data, master if master is not None else pd.DataFrame(), trust_clocks=True))
        if year is None or month is None:
            raise ValueError("Không đọc được tháng/năm từ file K9 (cần tên kiểu K9_07.2026.xlsx).")
        df = pd.DataFrame(rows)
        if log:
            log(f"Đã đọc {file.name}: {len(df)} dòng có giờ, kỳ {month:02d}/{year}.")
        df.attrs["roster"] = roster
        df.attrs["year"] = year
        df.attrs["month"] = month
        return df, int(year), int(month)
    finally:
        wb.close()


def load_working_table(path: str | Path, master: pd.DataFrame | None = None, log: LogFn = None) -> tuple[pd.DataFrame, int, int]:
    file = Path(path)
    raw = pd.read_excel(file, dtype=object)
    if raw.empty:
        raise ValueError(f"File trống: {file.name}")
    records = []
    for _, row in raw.iterrows():
        name = str(row.get("employee_name") or row.get("Employee Name") or "").strip()
        day = _as_work_date(row.get("date") or row.get("Date") or row.get("Ngày"))
        if not name or day is None:
            continue
        data = row.to_dict()
        data["employee_name"] = name
        data["date"] = day
        data["in1"] = _cell_time(data.get("in1"))
        data["out1"] = _cell_time(data.get("out1"))
        data["in2"] = _cell_time(data.get("in2"))
        data["out2"] = _cell_time(data.get("out2"))
        plus = max(
            clock_plus_days(row.get("in1")),
            clock_plus_days(row.get("out1")),
            clock_plus_days(row.get("in2")),
            clock_plus_days(row.get("out2")),
        )
        overnight = bool(data.get("overnight")) or plus > 0
        if isinstance(data.get("overnight"), str):
            overnight = data.get("overnight").strip().lower() in {"1", "true", "yes"} or plus > 0
        data["overnight"] = overnight
        data["manual_edit"] = bool(data.get("manual_edit")) or overnight
        records.append(recompute_record(data, master if master is not None else pd.DataFrame(), trust_clocks=True))
    df = pd.DataFrame(records)
    days = [_as_work_date(d) for d in df["date"].tolist()] if not df.empty else []
    days = [d for d in days if d]
    if not days:
        raise ValueError("Không có ngày hợp lệ trong file dữ liệu sửa giờ.")
    year, month = max(days).year, max(days).month
    counts: dict[tuple[int, int], int] = {}
    for d in days:
        counts[(d.year, d.month)] = counts.get((d.year, d.month), 0) + 1
    year, month = max(counts, key=counts.get)
    if log:
        log(f"Đã đọc {file.name}: {len(df)} dòng, kỳ {month:02d}/{year}.")
    df.attrs["year"] = year
    df.attrs["month"] = month
    return df, year, month


def load_attendance_source(
    path: str | Path, master: pd.DataFrame | None = None, log: LogFn = None
) -> tuple[pd.DataFrame, int, int]:
    file = Path(path)
    if not file.exists():
        raise FileNotFoundError(f"Không thấy file: {file}")
    if is_k9_workbook(file):
        return load_k9_file(file, master=master, log=log)
    return load_working_table(file, master=master, log=log)


def save_working_table(merged: pd.DataFrame, path: str | Path, log: LogFn = None) -> Path:
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if merged is None or merged.empty:
        pd.DataFrame(columns=WORKING_COLUMNS).to_excel(dest, index=False)
        return dest
    out = merged.copy()
    if "manual_edit" not in out.columns:
        out["manual_edit"] = False
    keep = [c for c in WORKING_COLUMNS if c in out.columns]
    extra = [c for c in out.columns if c not in keep]
    body = out[keep + [c for c in extra if c in ("name_key", "master_matched")]]
    body.to_excel(dest, index=False)
    if log:
        log(f"Đã ghi: {dest.name}  (file gốc để sửa giờ)")
    return dest


def days_in_table(df: pd.DataFrame) -> list[date]:
    if df is None or df.empty:
        return []
    days = sorted({d for d in (_as_work_date(v) for v in df["date"].tolist()) if d})
    return days


def annotate_day_rows(part: pd.DataFrame) -> pd.DataFrame:
    if part is None or part.empty:
        return part
    part = part.copy()
    part["_issue"] = part.apply(
        lambda r: ""
        if r.get("is_holiday")
        and not any((r.get("in1"), r.get("out1"), r.get("in2"), r.get("out2")))
        else (
            str(r.get("missing_punch") or "").strip()
            or issue_label(
                r.get("in1"),
                r.get("out1"),
                r.get("in2"),
                r.get("out2"),
                overnight=bool(r.get("overnight")),
                authorized_leave=_leave_flags(r)[0],
                half_day_leave=_leave_flags(r)[1],
            )
        ),
        axis=1,
    )
    marks = part.apply(milestone_from_row, axis=1)
    records = part.to_dict("records")
    holiday_rest = [
        bool(r.get("is_holiday")) and not any((r.get("in1"), r.get("out1"), r.get("in2"), r.get("out2")))
        for r in records
    ]
    leave_ok = [_leave_flags(r)[0] or _leave_flags(r)[1] for r in records]
    part["_reached"] = [
        True if rest or leave else m["reached"]
        for rest, leave, m in zip(holiday_rest, leave_ok, marks)
    ]
    part["_fraction"] = ["1.0 / lễ" if rest else m["fraction"] for rest, m in zip(holiday_rest, marks)]
    part["_badge"] = ["Ngày lễ" if rest else m["badge"] for rest, m in zip(holiday_rest, marks)]
    part["_cong_text"] = ["1.0" if rest else m["cong_text"] for rest, m in zip(holiday_rest, marks)]
    part["_hours_text"] = [m["hours_text"] for m in marks]
    part["_needs_review"] = [
        bool(issue) or (not reached)
        for issue, reached in zip(part["_issue"].astype(str), part["_reached"])
    ]
    return part


def _counts_from_annotated(part: pd.DataFrame) -> tuple[int, int, int]:
    if part is None or part.empty:
        return 0, 0, 0
    missing = int((part["_issue"].astype(str).str.len() > 0).sum())
    short = int((~part["_reached"].astype(bool)).sum())
    review = int(part["_needs_review"].astype(bool).sum())
    return review, missing, short


def day_review_count(df: pd.DataFrame, day: date) -> tuple[int, int, int]:
    """(cần xem, thiếu punch, chưa đạt 8h/12h)."""
    return _counts_from_annotated(rows_for_day(df, day, missing_only=False))


def review_counts_by_day(df: pd.DataFrame) -> dict[date, tuple[int, int, int]]:
    """Annotate once, then count per day (avoids 31× full scans)."""
    if df is None or df.empty:
        return {}
    part = annotate_day_rows(df)
    part = part.copy()
    part["_day"] = part["date"].map(_as_work_date)
    out: dict[date, tuple[int, int, int]] = {}
    for day, group in part.groupby("_day", sort=True):
        if day is None:
            continue
        out[day] = _counts_from_annotated(group)
    return out


def rows_for_day(df: pd.DataFrame, day: date, missing_only: bool = True) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()
    mask = df["date"].map(lambda v: _as_work_date(v) == day)
    part = annotate_day_rows(df.loc[mask])
    if part.empty:
        return part
    if missing_only:
        part = part[part["_issue"].astype(str).str.len() > 0]
        if "manual_edit" in part.columns:
            flagged = annotate_day_rows(df.loc[mask & df["manual_edit"].fillna(False)])
            part = pd.concat([part, flagged]).drop_duplicates(subset=["employee_name", "date"])
            part = annotate_day_rows(part)
    return sort_frame_by_employee_id(part)


def rows_for_employee(df: pd.DataFrame, name: str) -> pd.DataFrame:
    """All calendar days for one employee, annotated for the edit grid."""
    if df is None or df.empty or not str(name or "").strip():
        return pd.DataFrame()
    key = name_match_key(name)
    if not key:
        return pd.DataFrame()
    mask = df["employee_name"].map(lambda v: name_match_key(v) == key)
    part = annotate_day_rows(df.loc[mask])
    if part.empty:
        return part
    part = part.copy()
    part["_sort_day"] = part["date"].map(_as_work_date)
    return part.sort_values("_sort_day")


def upsert_clocks(
    df: pd.DataFrame,
    name: str,
    day: date,
    in1,
    out1,
    in2,
    out2,
    master: pd.DataFrame | None = None,
    manual: bool = True,
    overnight: bool | None = None,
) -> pd.DataFrame:
    """Insert or replace one employee-day using the 4 typed clocks, then recalc công."""
    master = master if master is not None else pd.DataFrame()
    key = name_match_key(name)
    in1, out1, in2, out2 = _as_time(in1), _as_time(out1), _as_time(in2), _as_time(out2)
    hit = None
    if df is not None and not df.empty:
        for idx, row in df.iterrows():
            if name_match_key(row.get("employee_name")) == key and _as_work_date(row.get("date")) == day:
                hit = idx
                break
    if hit is None:
        data = {
            "employee_name": name,
            "date": day,
            "in1": in1,
            "out1": out1,
            "in2": in2,
            "out2": out2,
            "overnight": bool(overnight),
            "manual_edit": manual,
        }
        data = recompute_record(data, master, trust_clocks=True)
        add = pd.DataFrame([data])
        return pd.concat([df, add], ignore_index=True) if df is not None and not df.empty else add
    data = df.loc[hit].to_dict()
    data["employee_name"] = name
    data["date"] = day
    data["in1"] = in1
    data["out1"] = out1
    data["in2"] = in2
    data["out2"] = out2
    if overnight is not None:
        data["overnight"] = bool(overnight)
    data["manual_edit"] = manual or bool(data.get("manual_edit"))
    data = recompute_record(data, master, trust_clocks=True)
    for col, value in data.items():
        df.at[hit, col] = value
    return df


def roster_names(df: pd.DataFrame, master: pd.DataFrame | None = None) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    id_by_key: dict[str, str] = {}

    def add(name: object, emp_id: object = "") -> None:
        text = str(name or "").strip()
        key = name_match_key(text)
        if not text or not key:
            return
        code = str(emp_id or "").strip()
        if code and key not in id_by_key:
            id_by_key[key] = code
        if key not in seen:
            seen.add(key)
            names.append(text)

    if df is not None and not df.empty and "employee_name" in df.columns:
        has_id = "employee_id" in df.columns
        for rec in df.itertuples(index=False):
            add(getattr(rec, "employee_name", ""), getattr(rec, "employee_id", "") if has_id else "")
    extra = list(df.attrs.get("roster") or []) if df is not None else []
    for name in extra:
        add(name)
    if master is not None and not master.empty and "employee_name" in master.columns:
        has_id = "employee_id" in master.columns
        for rec in master.itertuples(index=False):
            add(getattr(rec, "employee_name", ""), getattr(rec, "employee_id", "") if has_id else "")
    names.sort(key=lambda n: employee_id_sort_key(id_by_key.get(name_match_key(n), ""), n))
    return names


def reexport_corrected(
    df: pd.DataFrame,
    output_dir: str | Path,
    year: int,
    month: int,
    work_start: time = time(8, 0),
    log: LogFn = None,
) -> dict[str, Path]:
    from processing.cong_rules import filled_clocks
    from processing.reports import export_reports

    if df is None or df.empty:
        raise ValueError("Không có dữ liệu để xuất báo cáo.")
    has_clock = df.apply(
        lambda r: filled_clocks(r.get("in1"), r.get("out1"), r.get("in2"), r.get("out2")) > 0,
        axis=1,
    )
    df = df.loc[has_clock].copy()
    if df.empty:
        raise ValueError("Không còn dòng nào có giờ chấm sau khi lọc.")
    paths = export_reports(
        df,
        output_dir,
        work_start=work_start,
        log=log,
        year=year,
        month=month,
    )
    return paths


def _stamps_from_row(row) -> list:
    raw = row.get("punch_datetimes") if hasattr(row, "get") else None
    if isinstance(raw, (list, tuple)) and raw:
        return merge_day_punches(raw)
    from datetime import date as date_cls

    work = row.get("date") if hasattr(row, "get") else None
    if isinstance(work, datetime):
        work = work.date()
    if not isinstance(work, date_cls):
        work = _as_work_date(work)
    return merge_day_punches(
        clocks_to_datetimes(
            work,
            row.get("in1"),
            row.get("out1"),
            row.get("in2"),
            row.get("out2"),
            overnight=bool(row.get("overnight")),
        )
    )


def collect_missing_punches(df: pd.DataFrame) -> list[dict]:
    """Employee-days that still have exactly one punch after merge_day_punches."""
    if df is None or df.empty:
        return []
    out: list[dict] = []
    for _, rec in df.iterrows():
        full, half = _leave_flags(rec)
        if full:
            continue
        if rec.get("is_holiday") and not any((rec.get("in1"), rec.get("out1"), rec.get("in2"), rec.get("out2"))):
            continue
        stamps = _stamps_from_row(rec)
        status = str(rec.get("missing_punch") or "") or missing_punch_status(stamps)
        if not status or len(stamps) != 1:
            continue
        day = _as_work_date(rec.get("date"))
        recorded = stamps[0]
        out.append(
            {
                "employee_id": rec.get("employee_id") or "",
                "employee_name": rec.get("employee_name") or "",
                "date": day,
                "recorded": recorded,
                "recorded_text": recorded.strftime("%H:%M"),
                "status": status,
            }
        )
    out.sort(
        key=lambda item: (
            employee_id_sort_key(item.get("employee_id"), item.get("employee_name")),
            item.get("date") or date.min,
        )
    )
    return out


def apply_missing_punch_correction(
    df: pd.DataFrame,
    name: str,
    day: date,
    extra_hhmm: str,
    master: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Inject one typed clock, re-merge the day, and recalc công."""
    parsed = parse_time_value(extra_hhmm)
    if parsed is None:
        raise ValueError("Giờ bổ sung phải dạng HH:MM (ví dụ 17:30).")
    extra = datetime.combine(day, parsed)
    key = name_match_key(name)
    hit = None
    if df is not None and not df.empty:
        for idx, row in df.iterrows():
            if name_match_key(row.get("employee_name")) == key and _as_work_date(row.get("date")) == day:
                hit = idx
                break
    if hit is None:
        raise ValueError("Không tìm thấy dòng chấm công của nhân viên trong ngày này.")
    stamps = _stamps_from_row(df.loc[hit])
    merged = merge_day_punches(list(stamps) + [extra])
    if len(merged) < 2:
        raise ValueError("Giờ bổ sung trùng hoặc quá gần giờ đã ghi (< 5 phút). Nhập giờ khác.")
    in1, out1, in2, out2 = four_clocks_dt(merged)
    return upsert_clocks(
        df,
        name,
        day,
        in1.time() if in1 else None,
        out1.time() if out1 else None,
        in2.time() if in2 else None,
        out2.time() if out2 else None,
        master=master,
        manual=True,
    )


def clock_text(value, plus=0) -> str:
    from processing.sessions import format_clock_with_offset

    labeled = format_clock_with_offset(value, plus)
    return labeled or format_time(_as_time(value))

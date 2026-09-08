"""Detect the payroll month from the fingerprint source file, not from OCR/max date."""

from __future__ import annotations

import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import pandas as pd

# THÁNG 07.2026 / Thang 7-2026
_SHEET_PERIOD = re.compile(
    r"th[aáàảãạ]ng\s*(\d{1,2})\s*[./\-]\s*(\d{4})",
    re.IGNORECASE,
)
# 07.2026 or 07/2026 (K9 sheet / filename)
_MD_PERIOD = re.compile(r"(?<!\d)(\d{1,2})\s*[./]\s*(20\d{2})(?!\d)")
# 2026-07 or 2026_07
_YM_PERIOD = re.compile(r"(?<!\d)(20\d{2})\s*[-_]\s*(\d{1,2})(?!\d)")


def _valid(year: int, month: int) -> bool:
    return 2000 <= year <= 2100 and 1 <= month <= 12


def period_from_text(text: str | None) -> Optional[tuple[int, int]]:
    if not text:
        return None
    blob = str(text)
    match = _SHEET_PERIOD.search(blob)
    if match:
        month, year = int(match.group(1)), int(match.group(2))
        if _valid(year, month):
            return year, month
    match = _MD_PERIOD.search(blob)
    if match:
        month, year = int(match.group(1)), int(match.group(2))
        if _valid(year, month):
            return year, month
    match = _YM_PERIOD.search(blob)
    if match:
        year, month = int(match.group(1)), int(match.group(2))
        if _valid(year, month):
            return year, month
    return None


def _year_month(value: Any) -> Optional[tuple[int, int]]:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, datetime):
        return value.year, value.month
    if isinstance(value, date):
        return value.year, value.month
    if hasattr(value, "year") and hasattr(value, "month"):
        try:
            return int(value.year), int(value.month)
        except (TypeError, ValueError):
            return None
    return None


def period_from_dates(values) -> Optional[tuple[int, int]]:
    counts: Counter[tuple[int, int]] = Counter()
    for value in values:
        pair = _year_month(value)
        if pair:
            counts[pair] += 1
    if not counts:
        return None
    return counts.most_common(1)[0][0]


def detect_report_period(
    excel_path: str | Path | None = None,
    fingerprint_df: Optional[pd.DataFrame] = None,
    merged: Optional[pd.DataFrame] = None,
) -> tuple[int, int]:
    """Payroll month = fingerprint source (sheet/file/dates), never max(OCR date)."""
    if fingerprint_df is not None and not fingerprint_df.empty and "source_sheet" in fingerprint_df.columns:
        sheet_counts: Counter[tuple[int, int]] = Counter()
        for name in fingerprint_df["source_sheet"].dropna().astype(str):
            pair = period_from_text(name)
            if pair:
                sheet_counts[pair] += 1
        if sheet_counts:
            return sheet_counts.most_common(1)[0][0]

    if excel_path:
        path = Path(excel_path)
        pair = period_from_text(path.stem) or period_from_text(path.name)
        if pair:
            return pair
        if path.exists() and path.suffix.lower() in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook

                wb = load_workbook(path, read_only=True, data_only=True)
                try:
                    name_counts: Counter[tuple[int, int]] = Counter()
                    for sheet_name in wb.sheetnames:
                        found = period_from_text(sheet_name)
                        if found:
                            name_counts[found] += 1
                    if name_counts:
                        return name_counts.most_common(1)[0][0]
                finally:
                    wb.close()
            except Exception:  # noqa: BLE001 — fall through to dates
                pass

    if fingerprint_df is not None and not fingerprint_df.empty and "date" in fingerprint_df.columns:
        pair = period_from_dates(fingerprint_df["date"])
        if pair:
            return pair

    if merged is not None and not merged.empty and "date" in merged.columns:
        pair = period_from_dates(merged["date"])
        if pair:
            return pair

    today = date.today()
    return today.year, today.month


def filter_to_month(df: pd.DataFrame, year: int, month: int) -> pd.DataFrame:
    if df is None or df.empty or "date" not in df.columns:
        return df
    mask = df["date"].map(lambda value: _year_month(value) == (year, month))
    return df.loc[mask].copy()

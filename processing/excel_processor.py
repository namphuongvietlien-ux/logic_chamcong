"""Module 1: read and clean fingerprint-machine Excel exports."""

from __future__ import annotations

import re
from datetime import time
from pathlib import Path
from typing import Callable, Optional

import pandas as pd
from openpyxl import load_workbook

from processing.utils import max_time, min_time, parse_date_value, parse_time_value

LogFn = Optional[Callable[[str], None]]

EMPLOYEE_HEADER_RE = re.compile(
    r"Mã nhân viên:\s*(\S+)\s+Tên nhân viên:\s*(.+?)\s+Phòng ban:\s*(.*)",
    re.IGNORECASE,
)

GENERIC_NAME_COLS = {
    "tên nhân viên",
    "ho va ten",
    "họ và tên",
    "họ tên",
    "ho ten",
    "nhân viên",
    "nhan vien",
    "name",
    "employee",
    "employee name",
    "họ và tên nv",
}
GENERIC_DATE_COLS = {"ngày", "ngay", "date", "ngày công", "work date"}
GENERIC_IN_COLS = {"vào", "vao", "giờ vào", "gio vao", "check in", "check-in", "time in", "in"}
GENERIC_OUT_COLS = {"ra", "giờ ra", "gio ra", "check out", "check-out", "time out", "out"}
GENERIC_DT_COLS = {"thời gian", "thoi gian", "datetime", "punch time", "gio", "time"}


def _log(callback: LogFn, message: str) -> None:
    if callback:
        callback(message)


def _cell_str(value) -> str:
    return "" if value is None else str(value).strip()


def _looks_like_detail_header(value) -> bool:
    text = _cell_str(value)
    return "Tên nhân viên" in text and "Mã nhân viên" in text


def _classify_punches(times: list[time]) -> tuple[Optional[time], Optional[time]]:
    """Earliest punch = in, latest punch = out. Never invent a missing clock."""
    if not times:
        return None, None
    unique = sorted(set(times))
    if len(unique) == 1:
        only = unique[0]
        if only < time(12, 0):
            return only, None
        return None, only
    return unique[0], unique[-1]


def _parse_employee_header(text: str) -> tuple[str, str, str]:
    match = EMPLOYEE_HEADER_RE.search(text)
    if match:
        emp_id = match.group(1).strip()
        name = re.sub(r"\s+", " ", match.group(2)).strip(" -")
        dept = re.sub(r"\s+", " ", match.group(3)).strip(" -")
        if dept.replace("-", "") == "":
            dept = ""
        return emp_id, name, dept
    return "", re.sub(r"\s+", " ", text).strip(), ""


def load_fingerprint_blocks(path: Path, log: LogFn = None) -> pd.DataFrame:
    """Parse 'BẢNG CHI TIẾT CHẤM CÔNG' blocks used by CÔNG CHECK VÂN TAY THÁNG.xlsx."""
    rows: list[dict] = []
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            current_id = ""
            current_name = ""
            current_dept = ""
            in_table = False
            punch_cols: list[int] = []
            sheet_rows = 0
            for row in ws.iter_rows(values_only=True):
                values = list(row)
                if not values:
                    continue
                joined = " | ".join(_cell_str(v) for v in values if v is not None)
                header_cell = next((v for v in values if _looks_like_detail_header(v)), None)
                if header_cell is not None:
                    current_id, current_name, current_dept = _parse_employee_header(_cell_str(header_cell))
                    in_table = False
                    punch_cols = []
                    continue
                if current_name and any(_cell_str(v).upper() == "VÀO" for v in values):
                    punch_cols = [
                        idx
                        for idx, v in enumerate(values)
                        if _cell_str(v).upper() in {"VÀO", "RA", "VAO"}
                    ]
                    in_table = True
                    continue
                if current_name and joined.upper().startswith("TỔNG CỘNG"):
                    in_table = False
                    current_id, current_name, current_dept = "", "", ""
                    continue
                if not (current_name and in_table):
                    continue
                day = None
                for cell in values[:4]:
                    day = parse_date_value(cell)
                    if day:
                        break
                if day is None:
                    continue
                times: list[time] = []
                scan_indexes = punch_cols if punch_cols else list(range(3, min(8, len(values))))
                for idx in scan_indexes:
                    if idx >= len(values):
                        continue
                    parsed = parse_time_value(values[idx])
                    if parsed:
                        times.append(parsed)
                fp_in, fp_out = _classify_punches(times)
                if fp_in is None and fp_out is None:
                    continue
                rows.append(
                    {
                        "employee_id": current_id,
                        "employee_name": current_name,
                        "department": current_dept,
                        "date": day,
                        "fingerprint_in": fp_in,
                        "fingerprint_out": fp_out,
                        "punches": tuple(times),
                        "source_sheet": sheet_name,
                    }
                )
                sheet_rows += 1
            if sheet_rows:
                _log(log, f"Excel [{sheet_name}]: {sheet_rows} ngày chấm công.")
    finally:
        wb.close()
    return pd.DataFrame(rows)


def _normalize_header(value) -> str:
    return re.sub(r"\s+", " ", _cell_str(value)).strip().lower()


def load_fingerprint_tabular(path: Path, log: LogFn = None) -> pd.DataFrame:
    """Fallback: a flat table with name/date/in/out or a punch-log datetime column."""
    frames: list[pd.DataFrame] = []
    xls = pd.ExcelFile(path)
    for sheet in xls.sheet_names:
        raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=object)
        if raw.empty:
            continue
        header_row = None
        mapping: dict[str, int] = {}
        for idx, row in raw.head(30).iterrows():
            labels = [_normalize_header(v) for v in row.tolist()]
            name_i = next((i for i, lab in enumerate(labels) if lab in GENERIC_NAME_COLS), None)
            date_i = next((i for i, lab in enumerate(labels) if lab in GENERIC_DATE_COLS), None)
            in_i = next((i for i, lab in enumerate(labels) if lab in GENERIC_IN_COLS), None)
            out_i = next((i for i, lab in enumerate(labels) if lab in GENERIC_OUT_COLS), None)
            dt_i = next((i for i, lab in enumerate(labels) if lab in GENERIC_DT_COLS), None)
            if name_i is not None and (date_i is not None or dt_i is not None):
                header_row = int(idx)
                mapping = {"name": name_i, "date": date_i, "in": in_i, "out": out_i, "dt": dt_i}
                break
        if header_row is None:
            continue
        body = raw.iloc[header_row + 1 :]
        parsed_rows: list[dict] = []
        if mapping.get("dt") is not None and mapping.get("in") is None:
            punches: dict[tuple[str, object], list[time]] = {}
            names: dict[tuple[str, object], str] = {}
            for _, row in body.iterrows():
                name = _cell_str(row.iloc[mapping["name"]])
                if not name:
                    continue
                dt_cell = row.iloc[mapping["dt"]]
                day = parse_date_value(dt_cell)
                t = parse_time_value(dt_cell)
                if day is None or t is None:
                    continue
                key = (name, day)
                punches.setdefault(key, []).append(t)
                names[key] = name
            for (name, day), times in punches.items():
                fp_in, fp_out = _classify_punches(times)
                parsed_rows.append(
                    {
                        "employee_id": "",
                        "employee_name": name,
                        "date": day,
                        "fingerprint_in": fp_in,
                        "fingerprint_out": fp_out,
                        "punches": tuple(times),
                        "source_sheet": sheet,
                    }
                )
        else:
            for _, row in body.iterrows():
                name = _cell_str(row.iloc[mapping["name"]])
                if not name:
                    continue
                day = parse_date_value(row.iloc[mapping["date"]]) if mapping["date"] is not None else None
                if day is None:
                    continue
                fp_in = parse_time_value(row.iloc[mapping["in"]]) if mapping["in"] is not None else None
                fp_out = parse_time_value(row.iloc[mapping["out"]]) if mapping["out"] is not None else None
                if fp_in is None and fp_out is None:
                    continue
                parsed_rows.append(
                    {
                        "employee_id": "",
                        "employee_name": name,
                        "date": day,
                        "fingerprint_in": fp_in,
                        "fingerprint_out": fp_out,
                        "punches": tuple(t for t in (fp_in, fp_out) if t is not None),
                        "source_sheet": sheet,
                    }
                )
        if parsed_rows:
            frames.append(pd.DataFrame(parsed_rows))
            _log(log, f"Excel bảng phẳng [{sheet}]: {len(parsed_rows)} dòng.")
    if not frames:
        return pd.DataFrame()
    return pd.concat(frames, ignore_index=True)


def load_fingerprint_data(excel_path: str | Path, log: LogFn = None) -> pd.DataFrame:
    path = Path(excel_path)
    if not path.exists():
        raise FileNotFoundError(f"Không tìm thấy file Excel: {path}")
    _log(log, f"Đang đọc file vân tay: {path.name}")
    df = load_fingerprint_blocks(path, log=log)
    if df.empty:
        _log(log, "Không thấy khuôn BẢNG CHI TIẾT CHẤM CÔNG — thử đọc bảng phẳng.")
        df = load_fingerprint_tabular(path, log=log)
    if df.empty:
        raise ValueError(
            "Không đọc được dữ liệu chấm công từ Excel. "
            "Hãy chọn file 'CÔNG CHECK VÂN TAY THÁNG.xlsx' hoặc bảng có cột Tên / Ngày / Vào / Ra."
        )
    df["date"] = pd.to_datetime(df["date"]).dt.date
    df["employee_name"] = df["employee_name"].astype(str).str.replace(r"\s+", " ", regex=True).str.strip()
    if "punches" not in df.columns:
        df["punches"] = [() for _ in range(len(df))]
    if "employee_id" not in df.columns:
        df["employee_id"] = ""
    if "department" not in df.columns:
        df["department"] = ""

    def _best_punches(series):
        best: tuple = ()
        for val in series:
            if isinstance(val, (list, tuple)) and len(val) > len(best):
                best = tuple(val)
        return best

    df = (
        df.groupby(["employee_name", "date"], as_index=False)
        .agg(
            employee_id=("employee_id", "first"),
            department=("department", "first"),
            fingerprint_in=("fingerprint_in", lambda s: min_time(*list(s))),
            fingerprint_out=("fingerprint_out", lambda s: max_time(*list(s))),
            punches=("punches", _best_punches),
            source_sheet=("source_sheet", "first"),
        )
    )
    _log(log, f"Vân tay: {df['employee_name'].nunique()} nhân viên, {len(df)} ngày.")
    return df

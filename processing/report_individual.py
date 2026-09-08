"""Report 3: one formatted 'BẢNG CHI TIẾT CHẤM CÔNG' sheet per employee."""

from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from processing.excel_locale import excel_formula
from processing.leave import REMARK_PAID, REMARK_PAID_HALF, REMARK_UNPAID, REMARK_UNPAID_HALF
from processing.sessions import format_clock_with_offset
from processing.utils import format_date

LogFn = Optional[Callable[[str], None]]

YELLOW = PatternFill("solid", fgColor="FFFF00")
MISSING_FILL = PatternFill("solid", fgColor="FF9999")
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
TITLE_FONT = Font(name="Times New Roman", size=16, bold=True)
BOLD = Font(name="Times New Roman", size=11, bold=True)
NORMAL = Font(name="Times New Roman", size=11)
WEEKDAYS = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "CN"]


def _as_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _num(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    try:
        if pd.isna(value):
            return default
    except (TypeError, ValueError):
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", str(name or "NV")).strip() or "NV"
    cleaned = re.sub(r"\s+", " ", cleaned)[:31]
    title = cleaned
    n = 2
    while title.lower() in {x.lower() for x in used}:
        suffix = f"_{n}"
        title = (cleaned[: 31 - len(suffix)] + suffix)
        n += 1
    used.add(title)
    return title


def _border_range(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            cell.border = THIN
            if cell.alignment is None or not cell.alignment.horizontal:
                cell.alignment = CENTER


def _write_time(ws: Worksheet, row: int, col: int, value: Any, plus: int = 0) -> None:
    cell = ws.cell(row, col)
    cell.font = NORMAL
    cell.alignment = CENTER
    cell.border = THIN
    if value is None:
        cell.value = None
        return
    try:
        if pd.isna(value):
            cell.value = None
            return
    except (TypeError, ValueError):
        pass
    extra = 0
    try:
        extra = int(plus or 0)
    except (TypeError, ValueError):
        extra = 0
    if extra > 0:
        labeled = format_clock_with_offset(value, extra)
        cell.value = labeled
        return
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        cell.value = value
        cell.number_format = "HH:MM"
        return
    cell.value = str(value)


def _missing_flags(in1, out1, in2, out2, authorized_leave: bool = False, half_day: bool = False) -> list[bool]:
    """Highlight blank slots that should have a punch for this day."""
    slots = [in1, out1, in2, out2]
    present = [t is not None and not (isinstance(t, float) and pd.isna(t)) for t in slots]
    n = sum(1 for ok in present if ok)
    flags = [False, False, False, False]
    if authorized_leave:
        return flags
    if half_day:
        if n == 0:
            return flags
        if (present[0] and (present[1] or present[3])) or (present[2] and present[3]):
            return flags
    if n == 0:
        return flags
    if n == 1:
        if present[0] and not present[1] and not present[3]:
            flags[1] = True
        elif not present[0]:
            flags[0] = True
        return flags
    if n == 2:
        if not present[0]:
            flags[0] = True
        if not present[3] and not present[1]:
            flags[3] = True
        return flags
    if n == 3:
        for i, ok in enumerate(present):
            if not ok:
                flags[i] = True
    return flags


def _apply_sheet_layout(ws: Worksheet, header_end_row: int = 8) -> None:
    widths = {1: 14, 2: 12, 3: 10, 4: 10, 5: 10, 6: 10, 7: 10, 8: 10, 9: 10, 10: 10, 11: 10, 12: 11}
    for col in range(1, 16):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 12)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if header_end_row:
        ws.freeze_panes = f"A{header_end_row + 1}"
        ws.print_title_rows = f"1:{header_end_row}"


def _write_employee_sheet(
    ws: Worksheet,
    name: str,
    group: pd.DataFrame,
    start_row: int = 1,
    layout: bool = True,
) -> int:
    """Write one employee block starting at start_row. Returns the footer row index."""
    first = group.iloc[0]
    emp_id = str(first.get("employee_id") or "").strip() or "-----"
    dept = str(first.get("department") or "").strip() or "--------"
    s = int(start_row)

    total_hours = float(pd.to_numeric(group["actual_work_hours"], errors="coerce").fillna(0).sum())
    total_cong = float(pd.to_numeric(group["standardized_workday"], errors="coerce").fillna(0).sum())
    late_mins = int(pd.to_numeric(group["late_minutes"], errors="coerce").fillna(0).sum()) if "late_minutes" in group.columns else 0
    early_mins = int(pd.to_numeric(group["early_minutes"], errors="coerce").fillna(0).sum()) if "early_minutes" in group.columns else 0
    late_count = int((pd.to_numeric(group["late_minutes"], errors="coerce").fillna(0) > 0).sum()) if "late_minutes" in group.columns else 0
    early_count = int((pd.to_numeric(group["early_minutes"], errors="coerce").fillna(0) > 0).sum()) if "early_minutes" in group.columns else 0
    ot_hours = (
        float(pd.to_numeric(group["overtime_hours"], errors="coerce").fillna(0).sum())
        if "overtime_hours" in group.columns
        else 0.0
    )

    ws.merge_cells(start_row=s, start_column=1, end_row=s, end_column=15)
    title = ws.cell(s, 1, "BẢNG CHI TIẾT CHẤM CÔNG")
    title.font = TITLE_FONT
    title.alignment = CENTER
    ws.row_dimensions[s].height = 28

    ws.merge_cells(start_row=s + 1, start_column=1, end_row=s + 1, end_column=3)
    ws.merge_cells(start_row=s + 1, start_column=4, end_row=s + 1, end_column=10)
    ws.merge_cells(start_row=s + 1, start_column=11, end_row=s + 1, end_column=15)
    ws.cell(s + 1, 1, f"Mã nhân viên: {emp_id}")
    ws.cell(s + 1, 4, f"Tên nhân viên: {name}")
    ws.cell(s + 1, 11, f"Phòng ban: {dept}")
    for col in range(1, 16):
        cell = ws.cell(s + 1, col)
        cell.fill = YELLOW
        cell.font = BOLD
        cell.alignment = LEFT if col in (1, 4, 11) else CENTER
        cell.border = THIN
    ws.row_dimensions[s + 1].height = 22

    labels3 = [
        (1, "Tổng giờ"),
        (2, round(total_hours, 2)),
        (4, "Số lần trễ"),
        (5, late_count),
        (7, "Số phút trễ"),
        (8, late_mins),
    ]
    labels4 = [
        (1, "Tổng công"),
        (2, round(total_cong, 2)),
        (4, "Số lần sớm"),
        (5, early_count),
        (7, "Số phút sớm"),
        (8, early_mins),
    ]
    labels5 = [
        (1, "Tăng ca"),
        (2, round(ot_hours, 2)),
        (4, "Vắng KP"),
        (5, 0),
        (7, "Vắng CP"),
        (8, 0),
    ]
    for offset, pairs in ((2, labels3), (3, labels4), (4, labels5)):
        row_idx = s + offset
        for col, value in pairs:
            cell = ws.cell(row_idx, col, value)
            cell.font = BOLD if col in (1, 4, 7) else NORMAL
            cell.alignment = CENTER
            cell.border = THIN
        for col in range(1, 16):
            ws.cell(row_idx, col).border = THIN
            if ws.cell(row_idx, col).font is None:
                ws.cell(row_idx, col).font = NORMAL

    head = s + 6
    ws.merge_cells(start_row=head, start_column=3, end_row=head, end_column=4)
    ws.merge_cells(start_row=head, start_column=5, end_row=head, end_column=6)
    ws.merge_cells(start_row=head, start_column=7, end_row=head, end_column=15)
    ws.cell(head, 3, "1")
    ws.cell(head, 5, "2")
    ws.cell(head, 7, "Chi tiết")
    for col in range(1, 16):
        cell = ws.cell(head, col)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.fill = HEADER_FILL
        cell.border = THIN

    headers = ["Ngày", "Thứ", "Vào", "Ra", "Vào", "Ra", "Trễ", "Sớm", "Về trễ", "Giờ", "Công", "Tăng ca", "Ghi chú"]
    col_head = s + 7
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(col_head, col, text)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.fill = HEADER_FILL
        cell.border = THIN
    for col in range(13, 16):
        ws.cell(col_head, col).border = THIN
        ws.cell(col_head, col).fill = HEADER_FILL

    ordered = group.copy()
    ordered["_d"] = ordered["date"].map(_as_date)
    ordered = ordered.sort_values("_d")
    data_row = s + 8
    for _, rec in ordered.iterrows():
        day = _as_date(rec.get("date"))
        in1, out1, in2, out2 = rec.get("in1"), rec.get("out1"), rec.get("in2"), rec.get("out2")
        ws.cell(data_row, 1, day)
        ws.cell(data_row, 1).number_format = "DD/MM/YYYY"
        ws.cell(data_row, 1).font = NORMAL
        ws.cell(data_row, 1).alignment = CENTER
        ws.cell(data_row, 1).border = THIN
        ws.cell(data_row, 2, WEEKDAYS[day.weekday()] if day else "")
        ws.cell(data_row, 2).font = NORMAL
        ws.cell(data_row, 2).alignment = CENTER
        ws.cell(data_row, 2).border = THIN
        _write_time(ws, data_row, 3, in1, rec.get("in1_plus") or 0)
        _write_time(ws, data_row, 4, out1, rec.get("out1_plus") or 0)
        _write_time(ws, data_row, 5, in2, rec.get("in2_plus") or 0)
        _write_time(ws, data_row, 6, out2, rec.get("out2_plus") or 0)
        notes = str(rec.get("notes") or "")
        half_day = bool(rec.get("is_leave")) and float(rec.get("leave_duration") or 0) < 1
        half_day = half_day or REMARK_PAID_HALF in notes or REMARK_UNPAID_HALF in notes
        full_leave = bool(rec.get("is_leave")) and float(rec.get("leave_duration") or 0) >= 1
        full_leave = full_leave or (
            (REMARK_PAID in notes and REMARK_PAID_HALF not in notes)
            or (REMARK_UNPAID in notes and REMARK_UNPAID_HALF not in notes)
        )
        flags = _missing_flags(in1, out1, in2, out2, authorized_leave=full_leave, half_day=half_day)
        for offset, missing in enumerate(flags):
            if missing:
                ws.cell(data_row, 3 + offset).fill = MISSING_FILL
        for col, key, fmt in (
            (7, "late_minutes", "0"),
            (8, "early_minutes", "0"),
            (9, "late_return_minutes", "0"),
            (10, "actual_work_hours", "0.00"),
            (11, "standardized_workday", "0.0"),
            (12, "overtime_hours", "0.00"),
        ):
            val = rec.get(key)
            cell = ws.cell(data_row, col)
            cell.font = NORMAL
            cell.alignment = CENTER
            cell.border = THIN
            num = _num(val, 0)
            cell.value = num
            cell.number_format = fmt
        note_cell = ws.cell(data_row, 13, notes)
        note_cell.font = NORMAL
        note_cell.alignment = LEFT
        note_cell.border = THIN
        for col in range(14, 16):
            ws.cell(data_row, col).border = THIN
        data_row += 1

    last_data = max(data_row - 1, col_head)
    first_data = s + 8
    notes_range = f"M{first_data}:M{last_data}"
    cong_range = f"K{first_data}:K{last_data}"
    if last_data >= first_data:
        for dest_row, col_letter, fmt in (
            (s + 2, "J", "0.00"),
            (s + 3, "K", "0.00"),
            (s + 4, "L", "0.00"),
        ):
            cell = ws.cell(dest_row, 2, f"=SUM({col_letter}{first_data}:{col_letter}{last_data})")
            cell.number_format = fmt
            cell.font = NORMAL
            cell.alignment = CENTER
            cell.border = THIN
        authorized = "+".join(
            f'COUNTIF({notes_range}; "{tag}")'
            for tag in (REMARK_PAID, REMARK_PAID_HALF, REMARK_UNPAID, REMARK_UNPAID_HALF)
        )
        ws.cell(s + 4, 8, f"={authorized}")
        ws.cell(s + 4, 8).number_format = "0.0"
        ws.cell(s + 4, 8).font = NORMAL
        ws.cell(s + 4, 8).alignment = CENTER
        ws.cell(s + 4, 11, "Công nửa ngày phép")
        ws.cell(s + 4, 11).font = BOLD
        half_paid = ws.cell(
            s + 4, 12, excel_formula("SUMIF", notes_range, f'"{REMARK_PAID_HALF}"', cong_range)
        )
        half_paid.number_format = "0.00"
        half_paid.font = NORMAL
        half_paid.alignment = CENTER
    footer = last_data + 1
    ws.merge_cells(start_row=footer, start_column=1, end_row=footer, end_column=15)
    days_label = f"{total_cong:g}" if abs(total_cong - round(total_cong)) > 1e-6 else f"{int(round(total_cong))}"
    ws.cell(
        footer,
        1,
        f"TỔNG CỘNG: {days_label} NGÀY + {total_hours:.2f} GIỜ + {ot_hours:.2f} TĂNG CA",
    )
    ws.cell(footer, 1).font = BOLD
    ws.cell(footer, 1).alignment = CENTER
    ws.cell(footer, 1).fill = YELLOW
    for col in range(1, 16):
        ws.cell(footer, col).border = THIN
        ws.cell(footer, col).font = BOLD
        ws.cell(footer, col).fill = YELLOW

    _border_range(ws, head, last_data, 1, 15)
    if layout:
        _apply_sheet_layout(ws, col_head)
    return footer


def export_individual_detail(merged: pd.DataFrame, output_path: str | Path, log: LogFn = None) -> Path:
    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    default = wb.active
    used_names: set[str] = set()
    if merged is None or merged.empty:
        default.title = "Empty"
        default["A1"] = "Không có dữ liệu chấm công"
        wb.save(dest)
        wb.close()
        return dest

    names = sorted(str(n).strip() for n in merged["employee_name"].dropna().unique() if str(n).strip())
    combo_title = _safe_sheet_name("Tổng hợp", used_names)
    default.title = combo_title
    combo = default
    groups: list[tuple[str, pd.DataFrame]] = []
    for name in names:
        group = merged[merged["employee_name"].astype(str).str.strip() == name]
        if group.empty:
            continue
        groups.append((name, group))
        ws = wb.create_sheet(_safe_sheet_name(name, used_names))
        _write_employee_sheet(ws, name, group)

    row = 1
    for index, (name, group) in enumerate(groups):
        if index > 0:
            row += 1
        row = _write_employee_sheet(combo, name, group, start_row=row, layout=False) + 1
    _apply_sheet_layout(combo, header_end_row=0)
    combo.freeze_panes = "A9"
    combo.print_title_rows = "1:8"
    try:
        wb.move_sheet(combo, offset=-wb.sheetnames.index(combo.title))
    except (ValueError, TypeError):
        pass

    wb.save(dest)
    wb.close()
    if log:
        log(f"Đã ghi: {dest.name}  (Tổng hợp + {len(groups)} sheet nhân viên)")
    return dest
